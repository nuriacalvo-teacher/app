import json, re, sys
CARPETA = sys.argv[1] if len(sys.argv) > 1 else ''  # carpeta del archivo donde están estos expedientes
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def quitar_parentesis(txt):
    """Quita el primer grupo «( … )» con paréntesis anidados."""
    i = txt.find('(')
    if i < 0: return txt
    nivel = 0
    for j in range(i, len(txt)):
        if txt[j] == '(': nivel += 1
        elif txt[j] == ')':
            nivel -= 1
            if nivel == 0: return txt[:i] + txt[j+1:]
    return txt

filas = []
for f in ['resultado_doc1_a','resultado_doc1_b','resultado_doc2_a','resultado_doc2_b','resultado_doc3']:
    filas += json.load(open(f+'.json'))

NOTA_NO = 'Búsqueda web sin resultados relevantes (búsqueda por nombre completo, variantes y localidad).'
for r in filas:
    r['PROVINCIA'] = {'Lérida (Lleida)': 'Lleida', 'Lérida': 'Lleida'}.get(r['PROVINCIA'], r['PROVINCIA'])
    r['APELLIDOS'] = r['APELLIDOS'].replace('ABEILHE ', 'ABEILHÉ ')
    an = str(r.get('ANIO_NACIMIENTO') or '')
    m = re.match(r'(\d{4})', an)
    extra = 'Año de nacimiento ' + an[4:].strip(' ()') + '. ' if m and len(an) > 4 else ''
    r['NAC'] = m.group(1) if m else ''
    obs = r['OBSERVACIONES']
    if r['ILUSTRE'] == 'NO' and obs.startswith('Búsqueda sin resultados relevantes'):
        resto = obs[len('Búsqueda sin resultados relevantes'):].strip()
        resto = resto.lstrip(' ,')
        if resto.startswith('para '): resto = re.sub(r"^para ('[^']*'|[^.]*\.)\s*", '', resto)
        if resto.startswith('('): resto = quitar_parentesis(resto)
        resto = resto.lstrip(' .')
        obs = NOTA_NO + (' ' + resto if resto else '')
    # Sin referencias a páginas ni listas de «fuentes consultadas» que no se pudieron abrir
    obs = re.sub(r'\s*\((?:ver )?p\d+(?:, p\d+)?\)', '', obs)
    obs = re.sub(r';?\s*ver p\d+', '', obs)
    obs = re.sub(r'\s*Fuentes consultadas:.*?\.(?=\s|$)', '', obs)
    obs = re.sub(r',\s*p\d+\)', ')', obs)
    obs = obs.replace(' (p19-p22)', '')
    while '(fuentes consultadas' in obs:
        i = obs.index('(fuentes consultadas'); obs = obs[:i].rstrip() + quitar_parentesis(obs[i:])
    obs = re.sub(r'\s+([.,;])', r'\1', obs).replace('..', '.')
    r['OBS'] = (extra + obs).strip()
    if r['NOMBRE'] == 'Pascual' and r['APELLIDOS'].startswith('ABAD Y CASCAJARES'):
        r['NAC'] = ''
        r['OBS'] = ('Año de nacimiento dudoso: el extracto dice «de 14 años de edad» e ingreso el 26-IX-1883 (Colegio de Daroca); '
                    'nació hacia 1869 si la edad es la del ingreso. ') + r['OBSERVACIONES']

norm = lambda s: re.sub(r'[^A-ZÑ ]', '', s.upper().translate(str.maketrans('ÁÉÍÓÚÜ', 'AEIOUU')).replace(' Y ', ' '))
filas.sort(key=lambda r: (norm(r['APELLIDOS']), r['NOMBRE']))

wb = Workbook(); ws = wb.active; ws.title = 'Para revisar'
ws.append(['APELLIDOS', 'Nombre', 'País', 'Provincia', 'Localidad', 'Año nacimiento', 'Año expediente', 'Destacado o ilustre', 'Observaciones', 'Carpeta'])
colores = {'ILUSTRE': 'F3E7C9', 'DESTACADO': 'E3ECF6', 'SIN COMPROBAR': 'EEEEEE'}
for r in filas:
    ws.append([r['APELLIDOS'], r['NOMBRE'], r['PAIS'], r['PROVINCIA'], r['LOCALIDAD'], r['NAC'], str(r['ANIO_EXPEDIENTE']), r['ILUSTRE'], r['OBS'], CARPETA])
    c = ws.cell(ws.max_row, 8)
    if r['ILUSTRE'] in colores: c.fill = PatternFill('solid', fgColor=colores[r['ILUSTRE']]); c.font = Font(bold=r['ILUSTRE'] != 'SIN COMPROBAR')
fino = Side(style='thin', color='DCD6C8')
for i, a in enumerate([30, 20, 10, 14, 26, 11, 11, 16, 90, 10], 1): ws.column_dimensions[get_column_letter(i)].width = a
for fila in ws.iter_rows():
    for c in fila: c.alignment = Alignment(wrap_text=True, vertical='top'); c.border = Border(bottom=fino)
for c in ws[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F2A44')
for fila in ws.iter_rows(min_row=2, max_col=1): fila[0].font = Font(bold=True)
ws.freeze_panes = 'C2'; ws.auto_filter.ref = ws.dimensions
wb.save('../alumnos-xix/Prueba_expedientes_A.xlsx')
for r in filas:
    if r['ILUSTRE'] == 'NO': print(r['APELLIDOS'], '·', r['OBS'][:160])
